"""Utilities for extracting plain text from uploaded knowledge files."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


TEXT_EXTENSIONS = {".txt", ".md", ".text"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
TABLE_TEXT_EXTENSIONS = {".csv", ".tsv"}
WORD_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}

SUPPORTED_KNOWLEDGE_EXTENSIONS = (
    TEXT_EXTENSIONS
    | SPREADSHEET_EXTENSIONS
    | TABLE_TEXT_EXTENSIONS
    | WORD_EXTENSIONS
    | PDF_EXTENSIONS
)


def supported_knowledge_extensions_label() -> str:
    return ", ".join(sorted(SUPPORTED_KNOWLEDGE_EXTENSIONS))


def extract_knowledge_text(filename: str, raw: bytes) -> tuple[str, dict]:
    """Extract readable text and metadata from an uploaded knowledge file."""

    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_KNOWLEDGE_EXTENSIONS:
        raise ValueError(f"仅支持 {supported_knowledge_extensions_label()} 文件")

    if ext in TEXT_EXTENSIONS:
        content = _decode_text(raw)
        parser = "text"
    elif ext in TABLE_TEXT_EXTENSIONS:
        delimiter = "\t" if ext == ".tsv" else ","
        content = _extract_delimited_text(raw, delimiter)
        parser = "delimited"
    elif ext in WORD_EXTENSIONS:
        content = _extract_docx_text(raw)
        parser = "docx"
    elif ext in SPREADSHEET_EXTENSIONS:
        content = _extract_xlsx_text(raw)
        parser = "xlsx"
    elif ext in PDF_EXTENSIONS:
        content = _extract_pdf_text(raw)
        parser = "pdf"
    else:
        content = ""
        parser = "unknown"

    content = content.strip()
    if not content:
        raise ValueError("文件中没有提取到可用文本")

    return content, {"filename": filename, "file_ext": ext, "parser": parser}


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_delimited_text(raw: bytes, delimiter: str) -> str:
    text = _decode_text(raw)
    rows: list[str] = []
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for row in reader:
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_docx_text(raw: bytes) -> str:
    document = Document(io.BytesIO(raw))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_xlsx_text(raw: bytes) -> str:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        sheet_lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_to_text(cell) for cell in row]
            cells = [cell for cell in cells if cell]
            if cells:
                sheet_lines.append(" | ".join(cells))
        if sheet_lines:
            parts.append(f"## {sheet.title}\n" + "\n".join(sheet_lines))

    workbook.close()
    return "\n\n".join(parts)


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_pdf_text(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## 第 {index} 页\n{text}")
    return "\n\n".join(pages)
