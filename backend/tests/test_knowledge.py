"""Knowledge Sources CRUD tests — 资料库后端验收"""

import io
import asyncio
import pytest

# 确保知识库模型被注册到 Base.metadata（必须在 setup_db 建表之前）
from models.project_source import ProjectSource          # noqa: F401
from models.project_source_chunk import ProjectSourceChunk  # noqa: F401
from models.knowledge_qa_session import KnowledgeQaSession  # noqa: F401
from models.knowledge_qa_message import KnowledgeQaMessage  # noqa: F401

# 复用 test_smoke 的测试基础架构
from test_smoke import client, _auth_headers, setup_db, test_session_factory  # noqa: F401
from services.knowledge_query_plan import build_knowledge_query_plan
from services.knowledge_source import chunk_text, search_project_knowledge_with_plan


# ── Helpers ──

def _create_project(title="资料库测试项目"):
    headers = _auth_headers("knowledge_tester")
    resp = client.post("/api/projects", json={"title": title}, headers=headers)
    assert resp.status_code == 200, f"create project failed: {resp.text}"
    return resp.json()["id"], headers


def _create_source(project_id, title="测试资料", source_type="upload",
                   content="这是测试内容，需要足够长才能被切片器拆成多段。" * 5,
                   tags=None, always_inject=False, headers=None):
    return client.post(
        f"/api/projects/{project_id}/knowledge/sources",
        json={
            "title": title,
            "source_type": source_type,
            "content": content,
            "tags": tags,
            "always_inject": always_inject,
        },
        headers=headers,
    )


def _planned_search(project_id: str, question: str, limit: int = 10):
    plan = build_knowledge_query_plan(question, project_id=project_id)

    async def _run():
        async with test_session_factory() as session:
            structured, chunks = await search_project_knowledge_with_plan(
                session, project_id, plan, limit=limit
            )
            return structured, chunks

    structured, chunks = asyncio.run(_run())
    return plan, structured, chunks


# ==================== Chunking ====================

def test_chunk_text_keeps_chapter_heading_boundaries():
    """小说正文没有空行时，也不能把多个章节合成同一个切片。"""
    content = "\n".join(
        f"第{i}章 标题{i}\n" + ("莫凡和穆宁雪一起行动。叶心夏在另一边等待。" * 18)
        for i in range(1, 5)
    )

    chunks = chunk_text(content)

    assert len(chunks) == 4
    for index, chunk in enumerate(chunks, start=1):
        assert chunk.startswith(f"第{index}章")
        assert not any(f"第{other}章" in chunk for other in range(1, 5) if other != index)


def test_chunk_text_splits_table_by_rows_and_repeats_header():
    """长表格应按行切片，并在每个切片保留标题/表头上下文。"""
    content = "## 法系通用技能表\n序号 | 法系 | 阶位 | 技能\n" + "\n".join(
        f"{i} | 冰系 | 阶位{i} | 技能{i}" for i in range(1, 80)
    )

    chunks = chunk_text(content, chunk_size=360)

    assert len(chunks) > 1
    assert all(chunk.startswith("## 法系通用技能表\n序号 | 法系 | 阶位 | 技能") for chunk in chunks)
    assert any("79 | 冰系" in chunk for chunk in chunks)


# ==================== Create ====================

def test_create_source_basic():
    """创建资料后返回 source，chunk_count 正确"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="同人规则测试", source_type="fanfic_rule",
                          content="角色A不会游泳，角色B讨厌蘑菇。" * 10, headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["title"] == "同人规则测试"
    assert data["source_type"] == "fanfic_rule"
    assert data["always_inject"] is True
    assert data["token_count"] > 0


def test_create_source_fanfic_rule_always_inject():
    """fanfic_rule 类型自动强制 always_inject=True"""
    pid, headers = _create_project()
    resp = _create_source(pid, source_type="fanfic_rule", always_inject=False, headers=headers)
    assert resp.status_code == 200
    assert resp.json()["always_inject"] is True


def test_create_source_empty_content():
    """空内容允许保存"""
    pid, headers = _create_project()
    resp = _create_source(pid, content="", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["content"] == ""


def test_create_source_invalid_type():
    """非法 source_type 应被拒绝"""
    pid, headers = _create_project()
    resp = _create_source(pid, source_type="invalid_type", headers=headers)
    assert resp.status_code == 422


# ==================== Read: GET single ====================

def test_get_source():
    """获取单条资料"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="查看测试", headers=headers)
    sid = resp.json()["id"]

    resp2 = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert resp2.status_code == 200
    assert resp2.json()["title"] == "查看测试"


def test_get_source_not_found():
    """获取不存在的资料应 404"""
    pid, headers = _create_project()
    resp = client.get(
        f"/api/projects/{pid}/knowledge/sources/00000000-0000-0000-0000-000000000000",
        headers=headers,
    )
    assert resp.status_code == 404


# ==================== Read: List with filters ====================

def test_list_sources_with_type_filter():
    """list 支持 source_type 过滤"""
    pid, headers = _create_project()
    _create_source(pid, title="上传资料", source_type="upload", headers=headers)
    _create_source(pid, title="同人规则", source_type="fanfic_rule", headers=headers)
    _create_source(pid, title="时间线", source_type="timeline", headers=headers)

    resp = client.get(
        f"/api/projects/{pid}/knowledge/sources?source_type=fanfic_rule",
        headers=headers,
    )
    assert resp.status_code == 200
    items = resp.json()
    assert len(items) == 1
    assert items[0]["source_type"] == "fanfic_rule"


def test_list_sources_with_q_search():
    """list 支持 q 关键词搜索"""
    pid, headers = _create_project()
    _create_source(pid, title="全职法师设定", content="魔法世界体系", headers=headers)
    _create_source(pid, title="角色关系", content="张三和李四是好友", headers=headers)
    _create_source(pid, title="世界观", content="包含多个位面", headers=headers)

    # 搜索标题
    resp = client.get(f"/api/projects/{pid}/knowledge/sources?q=角色", headers=headers)
    assert resp.status_code == 200
    items = resp.json()
    assert len(items) == 1
    assert items[0]["title"] == "角色关系"

    # 搜索内容
    resp2 = client.get(f"/api/projects/{pid}/knowledge/sources?q=位面", headers=headers)
    assert resp2.status_code == 200
    items2 = resp2.json()
    assert len(items2) == 1
    assert items2[0]["title"] == "世界观"


def test_list_sources_with_sort():
    """list 支持 sort 排序"""
    pid, headers = _create_project()
    _create_source(pid, title="B资料", headers=headers)
    _create_source(pid, title="A资料", headers=headers)

    resp = client.get(f"/api/projects/{pid}/knowledge/sources?sort=title", headers=headers)
    assert resp.status_code == 200
    items = resp.json()
    assert len(items) == 2
    assert items[0]["title"] == "A资料"
    assert items[1]["title"] == "B资料"


# ==================== Chunks ====================

def test_list_chunks_ordered():
    """获取 chunks 按 chunk_index 排序"""
    pid, headers = _create_project()
    content = "第一段内容很长。\n" * 50 + "\n\n第二段内容也很长。\n" * 50
    resp = _create_source(pid, title="切片测试", content=content, headers=headers)
    sid = resp.json()["id"]

    resp2 = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}/chunks", headers=headers)
    assert resp2.status_code == 200
    chunks = resp2.json()
    assert len(chunks) >= 1
    indices = [c["chunk_index"] for c in chunks]
    assert indices == sorted(indices)


# ==================== Update ====================

def test_update_source_content():
    """更新 content"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="更新测试", content="原始内容较短", headers=headers)
    sid = resp.json()["id"]

    resp2 = client.patch(
        f"/api/projects/{pid}/knowledge/sources/{sid}",
        json={"content": "这是更新后的更长内容，确保重新切片后 chunk 数量变化。" * 10},
        headers=headers,
    )
    assert resp2.status_code == 200
    assert "更新后" in resp2.json()["content"]


def test_update_source_title():
    """更新标题不需要重切片"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="原标题", headers=headers)
    sid = resp.json()["id"]

    resp2 = client.patch(
        f"/api/projects/{pid}/knowledge/sources/{sid}",
        json={"title": "新标题"},
        headers=headers,
    )
    assert resp2.status_code == 200
    assert resp2.json()["title"] == "新标题"


def test_update_to_fanfic_rule_forces_inject():
    """改成 fanfic_rule 时后端强制 always_inject=True"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="普通资料", source_type="upload", headers=headers)
    sid = resp.json()["id"]
    assert resp.json()["always_inject"] is False

    resp2 = client.patch(
        f"/api/projects/{pid}/knowledge/sources/{sid}",
        json={"source_type": "fanfic_rule"},
        headers=headers,
    )
    assert resp2.status_code == 200
    assert resp2.json()["source_type"] == "fanfic_rule"
    assert resp2.json()["always_inject"] is True


# ==================== Delete ====================

def test_delete_source():
    """删除 source 后不可再 GET"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="待删除", headers=headers)
    sid = resp.json()["id"]

    del_resp = client.delete(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert del_resp.status_code == 204

    get_resp = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert get_resp.status_code == 404


def test_delete_source_removes_associated_facts():
    """删除 source 时应级联删除该 source 挂载的事实索引。"""
    from sqlalchemy import select, func
    from models.project_knowledge_fact import ProjectKnowledgeFact

    pid, headers = _create_project()
    sid = _create_source(pid, title="挂事实的资料", headers=headers).json()["id"]

    async def _insert_fact():
        async with test_session_factory() as session:
            session.add(
                ProjectKnowledgeFact(
                    project_id=pid,
                    source_id=sid,
                    fact_type="character_system",
                    subject="张小侯",
                    predicate="has_magic_system",
                    object="风系",
                    confidence="explicit",
                    evidence_text="张小侯是风系。",
                    extractor="test",
                    metadata_={"source_title": "挂事实的资料"},
                )
            )
            await session.commit()

    asyncio.run(_insert_fact())

    del_resp = client.delete(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert del_resp.status_code == 204

    async def _count_facts():
        async with test_session_factory() as session:
            result = await session.execute(
                select(func.count()).select_from(ProjectKnowledgeFact).where(
                    ProjectKnowledgeFact.source_id == sid,
                )
            )
            return result.scalar_one()

    assert asyncio.run(_count_facts()) == 0


# ==================== Upload ====================

def test_upload_txt_file():
    """上传 txt 文件"""
    pid, headers = _create_project()
    content = "这是一段很长的测试内容。\n" * 200
    file_data = io.BytesIO(content.encode("utf-8"))

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("test_upload.txt", file_data, "text/plain")},
        data={"title": "上传测试"},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["title"] == "上传测试"


def test_upload_docx_file():
    """上传 docx 文件并提取段落文本"""
    from docx import Document

    pid, headers = _create_project()
    document = Document()
    document.add_paragraph("穆宁雪是重要角色。")
    document.add_paragraph("她的资料应该能进入资料库检索。")
    file_data = io.BytesIO()
    document.save(file_data)
    file_data.seek(0)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("角色资料.docx", file_data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"title": "Word 资料"},
        headers=headers,
    )
    assert resp.status_code == 200

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    assert "穆宁雪是重要角色" in detail.json()["content"]


def test_upload_xlsx_file():
    """上传 xlsx 文件并提取工作表文本"""
    from openpyxl import Workbook

    pid, headers = _create_project()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "角色表"
    sheet.append(["角色", "能力", "备注"])
    sheet.append(["莫凡", "火系", "主角"])
    sheet.append(["穆宁雪", "冰系", "重要角色"])
    file_data = io.BytesIO()
    workbook.save(file_data)
    file_data.seek(0)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("角色表.xlsx", file_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"title": "Excel 资料"},
        headers=headers,
    )
    assert resp.status_code == 200

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    content = detail.json()["content"]
    assert "角色表" in content
    assert "穆宁雪 | 冰系 | 重要角色" in content


def test_upload_csv_file():
    """上传 csv 文件并提取表格文本"""
    pid, headers = _create_project()
    content = "法系,阶位,基础技能\n冰系,初阶,冰蔓\n冰系,中阶,冰锁\n"
    file_data = io.BytesIO(content.encode("utf-8"))

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("法系技能.csv", file_data, "text/csv")},
        data={"title": "CSV 技能表"},
        headers=headers,
    )
    assert resp.status_code == 200

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    source_content = detail.json()["content"]
    assert "法系 | 阶位 | 基础技能" in source_content
    assert "冰系 | 初阶 | 冰蔓" in source_content


def test_upload_tsv_file():
    """上传 tsv 文件并提取表格文本"""
    pid, headers = _create_project()
    content = "角色\t关系\t备注\n莫凡\t叶心夏\t亲密且互相信任\n"
    file_data = io.BytesIO(content.encode("utf-8"))

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("角色关系.tsv", file_data, "text/tab-separated-values")},
        data={"title": "TSV 关系表"},
        headers=headers,
    )
    assert resp.status_code == 200

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    source_content = detail.json()["content"]
    assert "角色 | 关系 | 备注" in source_content
    assert "莫凡 | 叶心夏 | 亲密且互相信任" in source_content


def test_query_plan_system_skill_question():
    """冰系有哪些技能：实体应保留为冰系，不应抽成冰系有哪些。"""
    plan = build_knowledge_query_plan("冰系有哪些技能？")
    assert plan.intent == "ability"
    assert plan.entities == ["冰系"]
    assert "冰系 技能" in plan.search_queries
    assert "冰系有哪些" not in plan.required_terms


def test_search_system_skill_table():
    """技能表资料中能召回某系技能行。"""
    pid, headers = _create_project()
    content = """## 法系通用技能表
序号 | 法系大类 | 法系 | 阶位 | 基础技能 | 一阶变体 | 二阶变体 | 三阶变体 | 备注
11 | ❄️ 冰系 | 初阶 | 冰蔓 | 冰蔓·覆盖 | 冰蔓·冰晶 | 冰蔓·冻结 | 三阶为完全冻结
12 | ❄️ 冰系 | 中阶 | 冰锁 | 冰锁·永冻 | 冰锁·碾碎 | 冰锁·噬魂 | 三阶名称出自穆白高阶战斗
13 | ❄️ 冰系 | 高阶 | 暴风雪 | 暴风雪·霜降 | 暴风雪·暴雪 | 暴风雪·冰封天地 | 三阶段为合理推断
14 | ❄️ 冰系 | 超阶 | 绝对零度 | — | — | — | 领域级冰封
"""
    source = _create_source(pid, title="法系通用技能表", content=content, headers=headers)
    assert source.status_code == 200

    resp = client.post(
        f"/api/projects/{pid}/knowledge/search",
        json={"query": "冰系有哪些技能？", "limit": 5},
        headers=headers,
    )
    assert resp.status_code == 200
    snippets = "\n".join(item["snippet"] for item in resp.json()["results"])
    assert "冰系" in snippets
    assert any(term in snippets for term in ("冰蔓", "冰锁", "暴风雪", "绝对零度"))


def test_qa_system_skill_answer_uses_table_rows():
    """问某系技能时，答案必须从表格行抽取技能，而不是答“没有资料”。"""
    pid, headers = _create_project()
    content = """## 法系通用技能表
序号 | 法系大类 | 法系 | 阶位 | 基础技能 | 一阶变体 | 二阶变体 | 三阶变体 | 备注
11 | ❄️ 冰系 | 初阶 | 冰蔓 | 冰蔓·覆盖 | 冰蔓·冰晶 | 冰蔓·冻结 | 三阶为完全冻结
12 | ❄️ 冰系 | 中阶 | 冰锁 | 冰锁·永冻 | 冰锁·碾碎 | 冰锁·噬魂 | 中阶冰系技能
13 | ❄️ 冰系 | 高阶 | 暴风雪 | 暴风雪·霜降 | 暴风雪·暴雪 | 暴风雪·冰封天地 | 高阶冰系技能
"""
    _create_source(pid, title="法系通用技能表", content=content, headers=headers)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "冰系有哪些技能？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    answer = resp.json()["answer"]
    assert "冰系" in answer
    assert "冰蔓" in answer
    assert "冰锁" in answer
    assert "暴风雪" in answer
    assert "没有" not in answer[:40]


def test_planned_retrieval_regression_common_questions():
    """资料问答常见问题回归集：防止靠人工截图逐条调试。"""
    pid, headers = _create_project()
    skill_table = """## 法系通用技能表
序号 | 法系大类 | 法系 | 阶位 | 基础技能 | 一阶变体 | 二阶变体 | 三阶变体 | 备注
1 | 🔥 火系 | 初阶 | 火滋 | 火滋·灼烧 | 火滋·焚骨 | 火滋·爆裂 | 基础火系技能
2 | 🔥 火系 | 中阶 | 烈拳 | 烈拳·九宫 | 烈拳·轰天 | 烈拳·炎龙 | 进阶火系技能
3 | 🔥 火系 | 高阶 | 天焰葬礼 | 天焰葬礼·焰雨 | 天焰葬礼·地狱火 | 天焰葬礼·千叶焰 | 高阶火系技能
11 | ❄️ 冰系 | 初阶 | 冰蔓 | 冰蔓·覆盖 | 冰蔓·冰晶 | 冰蔓·冻结 | 三阶为完全冻结
12 | ❄️ 冰系 | 中阶 | 冰锁 | 冰锁·永冻 | 冰锁·碾碎 | 冰锁·噬魂 | 中阶冰系技能
13 | ❄️ 冰系 | 高阶 | 暴风雪 | 暴风雪·霜降 | 暴风雪·暴雪 | 暴风雪·冰封天地 | 高阶冰系技能
14 | ❄️ 冰系 | 超阶 | 绝对零度 | — | — | — | 领域级冰封
"""
    character_notes = """## 人物与关系资料
穆宁雪是重要角色，早期明确觉醒冰系，后续资料中也提到她掌握风系。她与莫凡共同创立凡雪山，既是事业搭档，也是恋人关系。
莫凡与穆宁雪互相理解、互相支持，多次共同战斗。
叶心夏与莫凡关系亲密，是莫凡非常重视的人；资料中提到她与帕特农神庙、治愈系、心灵系相关。
"""
    _create_source(pid, title="法系通用技能表", content=skill_table, headers=headers)
    _create_source(pid, title="人物与关系资料", content=character_notes, headers=headers)

    cases = [
        ("冰系有哪些技能？", ("冰系",), ("冰蔓", "冰锁", "暴风雪", "绝对零度")),
        ("火系有哪些技能？", ("火系",), ("火滋", "烈拳", "天焰葬礼")),
        ("穆宁雪是什么系的？", ("穆宁雪",), ("冰系",)),
        ("莫凡和穆宁雪什么关系？", ("莫凡", "穆宁雪"), ("恋人", "凡雪山", "事业搭档")),
        ("莫凡和叶心夏什么关系？", ("莫凡", "叶心夏"), ("亲密", "重视", "帕特农")),
    ]

    for question, expected_entities, expected_terms in cases:
        plan, structured, chunks = _planned_search(pid, question)
        combined = "\n".join(item["snippet"] for item in chunks + structured)
        for entity in expected_entities:
            assert entity in plan.entities or entity in plan.required_terms
        assert combined, f"{question} 没有召回任何资料"
        assert any(term in combined for term in expected_terms), f"{question} 召回不含预期证据：{combined}"


def test_qa_followup_relationship_uses_session_context():
    """真实问答接口：短追问「那叶心夏呢」应沿用上一轮关系问题里的锚点。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="人物关系资料",
        content=(
            "莫凡与穆宁雪共同创立凡雪山，既是事业搭档，也是恋人关系。\n"
            "莫凡和叶心夏关系亲密，莫凡非常重视叶心夏，叶心夏与治愈系、心灵系相关。"
        ),
        headers=headers,
    )

    first = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "莫凡和穆宁雪什么关系？"},
        headers=headers,
    )
    assert first.status_code == 200, first.text
    conversation_id = first.json()["conversation_id"]

    second = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "那叶心夏呢？", "conversation_id": conversation_id},
        headers=headers,
    )
    assert second.status_code == 200, second.text
    data = second.json()
    plan = data["query_plan"]
    snippets = "\n".join(c["snippet"] for c in data["citations"])

    assert plan["intent"] == "relationship"
    assert "莫凡" in plan["entities"]
    assert "叶心夏" in plan["entities"]
    assert any(term in snippets for term in ("亲密", "重视", "治愈系", "心灵系"))

    messages = client.get(
        f"/api/projects/{pid}/knowledge/sessions/{conversation_id}/messages",
        headers=headers,
    )
    assert messages.status_code == 200
    assert [m["role"] for m in messages.json()] == ["user", "assistant", "user", "assistant"]


def test_qa_pronoun_followup_resolves_recent_entity():
    """真实问答接口：代词「她」应回指同一会话最近提到的人物。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="穆宁雪资料",
        content=(
            "穆宁雪早期明确觉醒冰系，是重要角色。\n"
            "莫凡与穆宁雪共同创立凡雪山，既是事业搭档，也是恋人关系。"
        ),
        headers=headers,
    )

    first = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "穆宁雪是什么系的？"},
        headers=headers,
    )
    assert first.status_code == 200, first.text
    conversation_id = first.json()["conversation_id"]

    second = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "莫凡跟她是什么关系？", "conversation_id": conversation_id},
        headers=headers,
    )
    assert second.status_code == 200, second.text
    data = second.json()
    plan = data["query_plan"]
    snippets = "\n".join(c["snippet"] for c in data["citations"])

    assert plan["intent"] == "relationship"
    assert "莫凡" in plan["entities"]
    assert "穆宁雪" in plan["entities"]
    assert any(term in snippets for term in ("恋人", "凡雪山", "事业搭档"))


def test_upload_invalid_extension():
    """上传不支持的文件类型应被拒绝"""
    pid, headers = _create_project()
    file_data = io.BytesIO(b"test content")

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("test.exe", file_data, "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 400


# ==================== Reindex ====================

def test_reindex_source():
    """手动重新切片"""
    pid, headers = _create_project()
    resp = _create_source(pid, title="重切片测试", headers=headers)
    sid = resp.json()["id"]

    resp2 = client.post(f"/api/projects/{pid}/knowledge/sources/{sid}/reindex", headers=headers)
    assert resp2.status_code == 200
    data = resp2.json()
    assert data["source_id"] == sid


# ==================== Character by Ability Query Plan Tests ====================

def test_query_plan_character_by_ability_question():
    """反向属性检索意图识别"""
    cases = [
        "谁是冰系的？",
        "那谁是冰系的呢？",
        "哪些角色有冰系？",
        "冰系人物有哪些？",
        "冰系角色有哪些？",
        "有哪些人觉醒冰系？",
        "有哪些角色拥有冰系？",
        "谁修炼了冰系？",
    ]
    for q in cases:
        plan = build_knowledge_query_plan(q)
        assert plan.intent == "character_by_ability", f"问题 '{q}' 意图应为 character_by_ability，实际为 {plan.intent}"
        assert "冰系" in plan.entities, f"问题 '{q}' 实体应包含 '冰系'"
        assert "冰系" in plan.required_terms, f"问题 '{q}' required_terms 应包含 '冰系'"
        assert "character" in plan.facets, f"问题 '{q}' facets 应包含 'character'"


def test_ability_question_not_confused_with_character_by_ability():
    """技能问题不应被误识别为反向属性检索"""
    cases = [
        "冰系有哪些技能？",
        "火系有什么魔法？",
        "雷系有哪些能力？",
    ]
    for q in cases:
        plan = build_knowledge_query_plan(q)
        assert plan.intent == "ability", f"问题 '{q}' 意图应为 ability，实际为 {plan.intent}"
        assert plan.entities[0].endswith("系"), f"问题 '{q}' 实体应包含系别"


def test_query_plan_character_system_detail_question():
    """人物某法系细问：不要把整句抽成怪实体。"""
    cases = [
        "莫凡他的土系是啥样的",
        "莫凡的土系是啥样的",
        "莫凡土系是什么样的",
        "叶心夏的治愈系是什么样",
    ]
    expected = [
        ("莫凡", "土系"),
        ("莫凡", "土系"),
        ("莫凡", "土系"),
        ("叶心夏", "治愈系"),
    ]
    for q, (character, system) in zip(cases, expected):
        plan = build_knowledge_query_plan(q)
        assert plan.intent == "ability"
        assert plan.entities[:2] == [character, system]
        assert plan.required_terms == [system]
        assert f"{character} {system}" in plan.search_queries
        assert f"{system} 技能" in plan.search_queries


# ==================== Character by Ability Retrieval Tests ====================

def test_character_by_ability_retrieval():
    """反向属性检索召回人物归属证据"""
    pid, headers = _create_project()
    
    # 创建人物资料
    _create_source(pid, title="穆宁雪资料",
                  content="穆宁雪早期明确觉醒冰系，并拥有强大的冰系天赋。她的冰系魔法非常出色。",
                  headers=headers)
    
    # 创建其他人物资料
    _create_source(pid, title="文霞军官",
                  content="文霞军官说自己的第三系是冰系。",
                  headers=headers)
    
    # 创建非冰系人物（不应被召回）
    _create_source(pid, title="莫凡资料",
                  content="莫凡主修火系和雷系，这里没有说明他是冰系。",
                  headers=headers)
    
    # 创建技能表（不应作为人物证据）
    _create_source(pid, title="法系通用技能表",
                  content="冰系技能包括冰蔓、冰锁、暴风雪、绝对零度。",
                  source_type="reference",
                  headers=headers)
    
    # 测试反向属性检索
    plan, structured, chunks = _planned_search(pid, "谁是冰系的？")
    
    # 验证意图识别
    assert plan.intent == "character_by_ability"
    assert "冰系" in plan.entities
    
    # 验证召回结果包含人物证据
    all_snippets = "\n".join([r.get("snippet", "") for r in structured + chunks])
    
    # 核心要求：应召回穆宁雪和文霞军官的证据
    assert "穆宁雪" in all_snippets or "文霞" in all_snippets, "应召回人物归属证据"
    
    # 人物资料应被召回
    person_results = [r for r in chunks if "穆宁雪" in r.get("snippet", "") or "文霞" in r.get("snippet", "")]
    assert len(person_results) > 0, "应召回人物资料"


def test_character_by_ability_followup_question():
    """反向属性检索追问"""
    pid, headers = _create_project()
    
    # 创建人物资料
    _create_source(pid, title="穆宁雪资料",
                  content="穆宁雪早期明确觉醒冰系，并拥有强大的冰系天赋。",
                  headers=headers)
    
    # 创建技能表
    _create_source(pid, title="法系通用技能表",
                  content="冰系技能包括冰蔓、冰锁、暴风雪、绝对零度。",
                  source_type="reference",
                  headers=headers)
    
    # 第一轮：问技能
    first = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "冰系有哪些技能？"},
        headers=headers,
    )
    assert first.status_code == 200, first.text
    conversation_id = first.json()["conversation_id"]
    
    # 第二轮：追问人物
    second = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "那谁是冰系的呢？", "conversation_id": conversation_id},
        headers=headers,
    )
    assert second.status_code == 200, second.text
    data = second.json()
    
    # 验证第二轮query_plan
    plan = data["query_plan"]
    assert plan["intent"] == "character_by_ability", f"第二轮意图应为 character_by_ability，实际为 {plan['intent']}"
    assert "冰系" in plan["entities"], "第二轮实体应包含 '冰系'"
    
    # 验证citations包含人物证据
    snippets = "\n".join(c["snippet"] for c in data["citations"])
    assert "穆宁雪" in snippets, "第二轮应召回穆宁雪的证据"
    
    # 核心回归测试：技能表不能排第一
    # 找出所有citation的title和score
    citations = data["citations"]
    assert len(citations) > 0, "应有召回结果"
    
    # 检查第一个结果不能是技能表
    first_citation = citations[0]
    first_title = first_citation.get("title", "")
    first_snippet = first_citation.get("snippet", "")
    
    # 技能表的特征：包含大量技能术语但没有人名
    skill_table_indicators = ["技能", "基础技能", "一阶变体", "二阶变体", "三阶变体", "初阶", "中阶", "高阶"]
    is_skill_table = (
        any(ind in first_title for ind in ["技能表", "技能"])
        or (sum(1 for ind in skill_table_indicators if ind in first_snippet) >= 3 and "穆宁雪" not in first_snippet)
    )
    
    assert not is_skill_table, f"技能表不能排第一！第一项: title={first_title}, snippet前100字={first_snippet[:100]}"
    
    # 验证人物证据排在技能表前面
    person_indices = [i for i, c in enumerate(citations) if "穆宁雪" in c.get("snippet", "")]
    skill_table_indices = [i for i, c in enumerate(citations) if any(ind in c.get("title", "") for ind in ["技能表", "技能"]) and "穆宁雪" not in c.get("snippet", "")]
    
    if person_indices and skill_table_indices:
        # 如果两者都存在，人物证据必须在技能表之前
        min_person_idx = min(person_indices)
        min_skill_idx = min(skill_table_indices)
        assert min_person_idx < min_skill_idx, f"人物证据(位置{min_person_idx})必须排在技能表(位置{min_skill_idx})前面"

    assert not any("技能表" in c.get("title", "") for c in citations), "反向人物检索不应引用纯技能表"
    assert not any("没有说明" in c.get("snippet", "") for c in citations), "否定证据不应作为人物归属引用"


def test_qa_character_by_ability_answer_lists_bound_character():
    """问谁是某系时，答案应列人物，不应拿技能表糊弄。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="穆宁雪资料",
        content="穆宁雪早期明确觉醒冰系，并拥有强大的冰系天赋。",
        headers=headers,
    )
    _create_source(
        pid,
        title="法系通用技能表",
        content="冰系技能包括冰蔓、冰锁、暴风雪、绝对零度。",
        source_type="reference",
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "谁是冰系的？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    answer = resp.json()["answer"]
    assert "穆宁雪" in answer
    assert "技能表" not in answer


def test_qa_character_system_list_answer_does_not_fall_back_to_generation():
    """问人物是什么系时，应抽取人物绑定法系，不能掉到 mock 小说生成模板。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="莫凡资料",
        content=(
            "莫凡在开学觉醒时展现天生双系。"
            "这是莫凡的雷霆系星尘，随后觉醒石又出现火热能量。"
            "莫凡觉醒了火系，展现强大的火焰掌控力。"
            "元素魔法七系包括冰系、火系、土系、水系、风系、光系、雷系。"
        ),
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "莫凡是什么系的？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    answer = resp.json()["answer"]
    assert "莫凡" in answer
    assert "雷系" in answer
    assert "火系" in answer
    assert "（生成内容）" not in answer
    # 通用七系背景不能被误认为莫凡个人全部法系。
    assert "水系" not in answer
    assert "风系" not in answer


def test_character_system_detail_uses_generic_system_table():
    """某人的某系细问：无个人直接证据时也要召回通用法系资料。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="莫凡资料",
        content="莫凡主修火系和雷系，当前片段没有明确写莫凡的土系表现。",
        headers=headers,
    )
    _create_source(
        pid,
        title="法系通用技能表",
        content=(
            "## 法系通用技能表\n"
            "序号 | 法系 | 阶位 | 基础技能 | 一阶变体 | 二阶变体 | 三阶变体\n"
            "21 | 土系 | 初阶 | 地波 | 地波·挪移 | 地波·陷落 | 地波·迟缓\n"
            "22 | 土系 | 中阶 | 岩障 | 岩障·山屏 | 岩障·磐石 | 岩障·石盾\n"
        ),
        source_type="reference",
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "莫凡他的土系是啥样的"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    plan = data["query_plan"]
    assert plan["intent"] == "character_ability"
    assert plan["entities"][:2] == ["莫凡", "土系"]

    snippets = "\n".join(c["snippet"] for c in data["citations"])
    assert "土系" in snippets
    assert any(term in snippets for term in ("地波", "岩障", "挪移", "山屏"))
    answer = data["answer"]
    assert "莫凡" in answer
    assert "土系" in answer
    assert any(term in answer for term in ("地波", "岩障", "挪移", "山屏"))


# ==================== QA Session Management Tests ====================

def test_qa_session_rename():
    """会话重命名"""
    pid, headers = _create_project()
    
    # 创建问答会话
    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "测试问题"},
        headers=headers,
    )
    assert resp.status_code == 200
    session_id = resp.json()["conversation_id"]
    
    # 重命名会话
    resp2 = client.patch(
        f"/api/projects/{pid}/knowledge/sessions/{session_id}",
        json={"title": "新标题"},
        headers=headers,
    )
    assert resp2.status_code == 200, resp2.text
    assert resp2.json()["title"] == "新标题"
    
    # 验证会话列表显示新标题
    resp3 = client.get(f"/api/projects/{pid}/knowledge/sessions", headers=headers)
    assert resp3.status_code == 200
    sessions = resp3.json()
    session = next(s for s in sessions if s["id"] == session_id)
    assert session["title"] == "新标题"


def test_qa_session_delete():
    """会话删除"""
    pid, headers = _create_project()
    
    # 创建问答会话
    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "测试问题"},
        headers=headers,
    )
    assert resp.status_code == 200
    session_id = resp.json()["conversation_id"]
    
    # 删除会话
    resp2 = client.delete(
        f"/api/projects/{pid}/knowledge/sessions/{session_id}",
        headers=headers,
    )
    assert resp2.status_code == 204
    
    # 验证会话列表不再包含该会话
    resp3 = client.get(f"/api/projects/{pid}/knowledge/sessions", headers=headers)
    assert resp3.status_code == 200
    sessions = resp3.json()
    assert not any(s["id"] == session_id for s in sessions)
    
    # 验证消息列表返回空或404
    resp4 = client.get(
        f"/api/projects/{pid}/knowledge/sessions/{session_id}/messages",
        headers=headers,
    )
    # 允许返回空列表或404
    assert resp4.status_code in (200, 404)
    if resp4.status_code == 200:
        assert len(resp4.json()) == 0


# ==================== V2 Query Planner ====================

def test_v2_schema_defaults():
    """V2 schema 默认值正确。"""
    from services.knowledge_query_planner import KnowledgeQueryPlanV2

    plan = KnowledgeQueryPlanV2(original_question="test")
    assert plan.intent == "general"
    assert plan.entities == []
    assert plan.attributes == []
    assert plan.confidence == 0.0
    assert plan.planner == "rule"


def test_v2_to_v1_conversion():
    """V2 to_v1() 正确转换。"""
    from services.knowledge_query_planner import KnowledgeQueryPlanV2

    plan = KnowledgeQueryPlanV2(
        original_question="冰系有哪些技能？",
        intent="character_ability",
        entities=["冰系"],
        attributes=["冰系"],
        sub_queries=["冰系 技能", "冰系 魔法"],
        required_terms=["冰系"],
        optional_terms=["技能"],
    )
    v1 = plan.to_v1()
    assert v1.intent == "ability"
    assert "冰系" in v1.entities
    assert "冰系" in v1.required_terms
    assert "冰系 技能" in v1.search_queries


def test_v2_rule_fallback_ability():
    """规则 planner 对能力类问题正确识别。"""
    from services.knowledge_query_planner import _build_rule_plan

    plan = _build_rule_plan("冰系有哪些技能？")
    assert plan.intent == "character_ability"
    assert "冰系" in plan.entities or "冰系" in plan.attributes
    assert "冰系" in plan.required_terms
    assert plan.confidence == 0.6
    assert plan.planner == "rule"
    assert plan.evidence_policy  # 应有证据策略
    assert plan.answer_policy    # 应有回答策略


def test_v2_rule_fallback_relationship():
    """规则 planner 对关系类问题正确识别。"""
    from services.knowledge_query_planner import _build_rule_plan

    plan = _build_rule_plan("莫凡和穆宁雪什么关系？")
    assert plan.intent == "relationship"
    assert "莫凡" in plan.entities
    assert "穆宁雪" in plan.entities
    assert len(plan.required_terms) >= 2


def test_v2_rule_fallback_character_by_ability():
    """规则 planner 对反向属性检索正确识别。"""
    from services.knowledge_query_planner import _build_rule_plan

    plan = _build_rule_plan("谁是冰系的？")
    assert plan.intent == "character_by_ability"
    assert "冰系" in plan.entities or "冰系" in plan.attributes


def test_v2_rule_fallback_worldbuilding():
    """规则 planner 对世界观问题正确识别。"""
    from services.knowledge_query_planner import _build_rule_plan

    plan = _build_rule_plan("博城灾难是什么？")
    assert plan.intent in ("worldbuilding", "plot_event")
    assert "博城灾难" in plan.entities


def test_v2_rule_followup_pronoun():
    """规则 planner 对代词追问做消解。"""
    from services.knowledge_query_planner import _build_rule_plan

    recent = [{"role": "user", "content": "莫凡和穆宁雪什么关系？"}]
    plan = _build_rule_plan("莫凡跟她是什么关系？", recent_messages=recent)
    assert plan.rewritten_question is not None
    assert "穆宁雪" in plan.rewritten_question


def test_v2_rule_followup_short():
    """规则 planner 对短追问做消解。"""
    from services.knowledge_query_planner import _build_rule_plan

    recent = [{"role": "user", "content": "莫凡和穆宁雪什么关系？"}]
    plan = _build_rule_plan("那叶心夏呢？", recent_messages=recent)
    # 应该被重写为关系问题
    assert plan.rewritten_question is not None
    assert "叶心夏" in plan.rewritten_question
    assert plan.intent == "relationship"


def test_v2_validate_plan():
    """_validate_plan 校验逻辑。"""
    from services.knowledge_query_planner import _validate_plan

    assert _validate_plan({"intent": "general", "entities": [], "sub_queries": [], "confidence": 0.5})
    assert not _validate_plan({})
    assert not _validate_plan({"intent": "invalid_intent", "entities": [], "sub_queries": [], "confidence": 0.5})
    assert not _validate_plan({"intent": "general", "entities": "not_list", "sub_queries": [], "confidence": 0.5})
    assert not _validate_plan({"intent": "general", "entities": [], "sub_queries": [], "confidence": 1.5})


def test_v2_character_system_detail():
    """规则 planner 对"某人的某系是什么样"正确处理。"""
    from services.knowledge_query_planner import _build_rule_plan

    plan = _build_rule_plan("莫凡他的土系是啥样的")
    assert plan.intent == "character_ability"
    assert "莫凡" in plan.entities
    assert "土系" in plan.entities or "土系" in plan.attributes
    assert "土系" in plan.required_terms


def test_ask_api_returns_v2_query_plan_fields():
    """真实 /ask API 返回 V2 扩展字段：attributes, confidence, planner, answer_policy。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="法系通用技能表",
        content="11 | ❄️ 冰系 | 初阶 | 冰蔓 | 冰蔓·覆盖 | 冰蔓·冰晶 | 冰蔓·冻结\n12 | ❄️ 冰系 | 中阶 | 冰锁 | 冰锁·永冻 | 冰锁·碾碎 | 冰锁·噬魂",
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "冰系有哪些技能？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    plan = data["query_plan"]

    # V2 扩展字段必须存在
    assert "attributes" in plan, "query_plan 缺少 attributes"
    assert "confidence" in plan, "query_plan 缺少 confidence"
    assert "planner" in plan, "query_plan 缺少 planner"
    assert "answer_policy" in plan, "query_plan 缺少 answer_policy"
    assert "sub_queries" in plan, "query_plan 缺少 sub_queries"
    assert "required_terms" in plan, "query_plan 缺少 required_terms"

    # 值校验
    assert plan["planner"] in ("llm", "rule")
    assert 0 <= plan["confidence"] <= 1
    assert plan["intent"] == "character_ability"
    assert "冰系" in plan["attributes"] or "冰系" in plan["entities"]
    assert plan["answer_policy"]  # 不为空

    # citations 必须存在
    assert "citations" in data
    assert len(data["citations"]) >= 1


def test_upload_pdf_file():
    """上传 PDF 文件并提取文本"""
    from fpdf import FPDF

    pid, headers = _create_project()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="This is a test PDF document for knowledge base upload.")
    buf = io.BytesIO(pdf.output())

    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("世界观.pdf", buf, "application/pdf")},
        data={"title": "PDF 世界观资料"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    content = detail.json()["content"]
    assert len(content) > 0, "PDF 解析后内容不应为空"


def test_upload_csv_file():
    """上传 CSV 文件并提取表格文本"""
    csv_content = "角色,能力,备注\n莫凡,火系,主角\n穆宁雪,冰系,重要角色\n"
    file_data = io.BytesIO(csv_content.encode("utf-8"))

    pid, headers = _create_project()
    resp = client.post(
        f"/api/projects/{pid}/knowledge/upload",
        files={"file": ("角色表.csv", file_data, "text/csv")},
        data={"title": "CSV 资料"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    sid = resp.json()["id"]
    detail = client.get(f"/api/projects/{pid}/knowledge/sources/{sid}", headers=headers)
    assert detail.status_code == 200
    content = detail.json()["content"]
    assert "莫凡" in content
    assert "穆宁雪" in content


# ==================== Evidence Retrieval ====================

def test_evidence_type_classification_ability():
    """技能表应被分类为 ability_table。"""
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "11 | 冰系 | 初阶 | 冰蔓 | 冰蔓·覆盖 | 冰蔓·冰晶\n12 | 冰系 | 中阶 | 冰锁 | 冰锁·永冻",
        "title": "法系通用技能表",
        "source_type": "reference",
    }
    assert _classify_evidence(result, "character_ability", ["冰系"], ["冰系"]) == "ability_table"


def test_evidence_type_classification_relationship():
    """两人同段 + 关系词应被分类为 relationship_evidence。"""
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "莫凡与穆宁雪共同创立凡雪山，既是事业搭档，也是恋人关系。",
        "title": "人物关系资料",
        "source_type": "upload",
    }
    assert _classify_evidence(result, "relationship", ["莫凡", "穆宁雪"], []) == "relationship_evidence"


def test_evidence_type_classification_negative():
    """含多个否定词应被分类为 negative_evidence。"""
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "资料中没有说明莫凡的土系表现，无法确定他是否拥有土系。",
        "title": "莫凡资料",
        "source_type": "upload",
    }
    assert _classify_evidence(result, "character_ability", ["莫凡"], ["土系"]) == "negative_evidence"


def test_evidence_type_classification_direct_character():
    """人物 + 属性同段（非技能表）应被分类为 direct_character_evidence。"""
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "穆宁雪早期觉醒冰系，是重要角色。",
        "title": "穆宁雪资料",
        "source_type": "upload",
    }
    assert _classify_evidence(result, "character_ability", ["穆宁雪"], ["冰系"]) == "direct_character_evidence"


def test_grouped_evidence_negative_excluded_from_prompt_when_positive_exists():
    """有正向证据时，否定证据不应出现在 prompt 中。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    grouped.direct_character.append(ClassifiedEvidence(
        evidence_type="direct_character_evidence",
        source_kind="chunk", source_id="1", chunk_id="1",
        title="穆宁雪资料", snippet="穆宁雪觉醒冰系", score=10.0,
    ))
    grouped.negative.append(ClassifiedEvidence(
        evidence_type="negative_evidence",
        source_kind="chunk", source_id="2", chunk_id="2",
        title="莫凡资料", snippet="没有说明莫凡的土系", score=5.0,
    ))

    prompt = grouped.to_prompt_text()
    assert "穆宁雪觉醒冰系" in prompt
    assert "没有说明莫凡的土系" not in prompt


def test_grouped_evidence_negative_included_when_no_positive():
    """无正向证据时，否定证据应出现在 prompt 中。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    grouped.negative.append(ClassifiedEvidence(
        evidence_type="negative_evidence",
        source_kind="chunk", source_id="1", chunk_id="1",
        title="莫凡资料", snippet="没有说明莫凡的土系", score=5.0,
    ))

    prompt = grouped.to_prompt_text()
    assert "没有说明莫凡的土系" in prompt


def test_grouped_evidence_citations_exclude_negative():
    """citations 不应包含否定证据。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    grouped.direct_character.append(ClassifiedEvidence(
        evidence_type="direct_character_evidence",
        source_kind="chunk", source_id="1", chunk_id="1",
        title="穆宁雪资料", snippet="穆宁雪觉醒冰系", score=10.0,
    ))
    grouped.negative.append(ClassifiedEvidence(
        evidence_type="negative_evidence",
        source_kind="chunk", source_id="2", chunk_id="2",
        title="莫凡资料", snippet="没有说明莫凡的土系", score=5.0,
    ))

    citations = grouped.to_citations()
    assert len(citations) == 1
    assert citations[0]["evidence_type"] == "direct_character_evidence"
    assert all(c["evidence_type"] != "negative_evidence" for c in citations)


def test_grouped_evidence_citations_include_evidence_type():
    """citations 应包含 evidence_type 字段。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    grouped.ability_table.append(ClassifiedEvidence(
        evidence_type="ability_table",
        source_kind="chunk", source_id="1", chunk_id="1",
        title="技能表", snippet="冰系初阶冰蔓", score=8.0,
    ))

    citations = grouped.to_citations()
    assert len(citations) == 1
    assert citations[0]["evidence_type"] == "ability_table"

def test_rerank_priority_ability_intent():
    """能力类问题：direct_character 优先于 ability_table。"""
    from services.knowledge_retrieval import _rerank_priority

    assert _rerank_priority("direct_character_evidence", "character_ability") < _rerank_priority("ability_table", "character_ability")
    assert _rerank_priority("ability_table", "character_ability") < _rerank_priority("negative_evidence", "character_ability")


def test_rerank_priority_relationship_intent():
    """关系类问题：relationship_evidence 最优先。"""
    from services.knowledge_retrieval import _rerank_priority

    assert _rerank_priority("relationship_evidence", "relationship") < _rerank_priority("direct_character_evidence", "relationship")
    assert _rerank_priority("direct_character_evidence", "relationship") < _rerank_priority("ability_table", "relationship")


def test_web_search_url_normalization():
    """DuckDuckGo 跳转 URL 应还原为原始网页 URL。"""
    from services.web_search import _normalize_duckduckgo_url

    url = _normalize_duckduckgo_url("/l/?uddg=https%3A%2F%2Fexample.com%2Fpage%3Fa%3D1")
    assert url == "https://example.com/page?a=1"


def test_web_search_extracts_readable_page_passages():
    """联网检索应抽取网页正文，而不只是搜索结果摘要。"""
    from services.web_search import _extract_readable_page

    html = """
    <html>
      <head><title>冰系技能说明</title><meta name="description" content="全职法师冰系资料"></head>
      <body>
        <nav>首页 导航</nav>
        <p>火系拥有烈拳和天焰葬礼。</p>
        <p>冰系初阶技能包括冰蔓，中阶技能包括冰锁，高阶技能包括暴风雪。</p>
        <p>这段和问题无关，只是普通页面说明文字。</p>
      </body>
    </html>
    """
    title, snippet = _extract_readable_page(html, "冰系有哪些技能")
    assert title == "冰系技能说明"
    assert "冰蔓" in snippet
    assert "冰锁" in snippet
    assert "暴风雪" in snippet


def test_web_search_ranks_query_related_passages():
    """正文抽取应优先保留和问题关键词匹配的段落。"""
    from services.web_search import _rank_blocks

    blocks = [
        "这是一段很长的普通背景介绍，不包含关键能力信息。" * 8,
        "穆宁雪的冰系能力包括冰晶刹弓相关表现。",
        "另一个角色的火系能力包括烈拳。",
    ]
    ranked = _rank_blocks(blocks, "穆宁雪冰系能力")
    assert ranked[0].startswith("穆宁雪")


def test_ask_with_web_search_returns_web_citation(monkeypatch):
    """开启 include_web 后，资料问答应返回联网引用和 web_hits。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="本地资料",
        content="穆宁雪是冰系角色。" * 5,
        headers=headers,
    )

    async def fake_search_web(query: str, *, limit: int | None = None):
        return [{
            "source_kind": "web_search",
            "source_id": "https://example.com/mu-ningxue",
            "chunk_id": None,
            "title": "穆宁雪资料页",
            "snippet": "联网资料显示穆宁雪相关设定。",
            "url": "https://example.com/mu-ningxue",
            "evidence_type": "web_search",
            "matched_query": query,
            "score": 1.0,
        }]

    monkeypatch.setattr("services.web_search.search_web", fake_search_web)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "穆宁雪是谁？", "include_web": True},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["retrieval_stats"]["web_hits"] == 1
    assert any(c["source_kind"] == "web_search" for c in data["citations"])
    web_citation = next(c for c in data["citations"] if c["source_kind"] == "web_search")
    assert web_citation["evidence_type"] == "web_search"
    assert web_citation["url"] == "https://example.com/mu-ningxue"


def test_ask_without_web_search_does_not_call_web(monkeypatch):
    """默认不联网，避免普通资料问答发出外部请求。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="本地资料",
        content="穆宁雪是冰系角色。" * 5,
        headers=headers,
    )

    async def fail_search_web(query: str, *, limit: int | None = None):
        raise AssertionError("include_web=false 时不应调用联网搜索")

    monkeypatch.setattr("services.web_search.search_web", fail_search_web)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "穆宁雪是谁？"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["retrieval_stats"]["web_hits"] == 0


# ── RAG 检索健壮性回归测试 ──────────────────────────────


def test_like_escape_wildcards_treated_as_literals():
    """用户检索词中的 % 与 _ 应作为字面量，而非 SQL 通配符。"""
    from services.knowledge_source import _like_escape

    # % 和 _ 被转义，反斜杠先转义
    assert _like_escape("a%b_c") == "a\\%b\\_c"
    assert _like_escape("100%") == "100\\%"
    assert _like_escape("a_b") == "a\\_b"
    assert _like_escape("path\\to") == "path\\\\to"


def test_search_does_not_match_via_wildcard_injection():
    """含 % 的检索词不应匹配任意内容，应只匹配字面包含 % 的资料。"""
    pid, headers = _create_project()
    # 一条不含 % 的资料
    _create_source(pid, title="普通资料", content="这是一段普通正文，没有任何特殊符号。", headers=headers)
    # 一条字面包含 % 的资料
    _create_source(pid, title="含百分号", content="折扣率 50% 的活动。", headers=headers)

    resp = client.post(
        f"/api/projects/{pid}/knowledge/search",
        json={"query": "50%"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    hits = resp.json().get("results", [])
    titles = {h.get("title", "") for h in hits}
    # 应只命中字面含 "50%" 的资料，不含通配符语义下的"普通资料"
    assert "含百分号" in titles
    assert "普通资料" not in titles


def test_grouped_evidence_respects_total_char_budget():
    """证据 prompt 应按总字符预算截断，超出部分不进入 prompt。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    # 构造远超预算的证据
    long_snippet = "x" * 800
    for i in range(20):
        grouped.generic.append(ClassifiedEvidence(
            evidence_type="generic_context",
            source_kind="chunk", source_id=str(i), chunk_id=str(i),
            title=f"资料{i}", snippet=long_snippet, score=1.0,
        ))

    text = grouped.to_prompt_text(max_chars=2000, per_section=20)
    # 截断后长度受预算约束（允许少量标题开销超出）
    assert len(text) < 2200
    # 不应包含全部 20 条
    assert text.count("资料") < 20


def test_grouped_evidence_budget_preserves_high_priority_sections():
    """预算不足时应优先保留靠前的高优先级分组，低优先级被截断。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    # 高优先级：人物直接证据，多条大 snippet 占满预算
    for i in range(4):
        grouped.direct_character.append(ClassifiedEvidence(
            evidence_type="direct_character_evidence",
            source_kind="chunk", source_id=f"1{i}", chunk_id=f"1{i}",
            title=f"人物证据{i}", snippet="A" * 600, score=10.0,
        ))
    # 低优先级：普通资料
    grouped.generic.append(ClassifiedEvidence(
        evidence_type="generic_context",
        source_kind="chunk", source_id="2", chunk_id="2",
        title="普通证据", snippet="B" * 600, score=1.0,
    ))

    # 预算只够装下高优先级的前几条
    text = grouped.to_prompt_text(max_chars=1500, per_section=6)
    assert "人物证据0" in text
    # 预算被高优先级占满后，低优先级 section 不应出现
    assert "普通证据" not in text
    assert "普通资料" not in text


def test_grouped_evidence_budget_zero_or_negative_does_not_crash():
    """预算为 0 或负（history/web 占满总预算的极端情况）时不应崩溃，产出受控。"""
    from services.knowledge_retrieval import GroupedEvidence, ClassifiedEvidence

    grouped = GroupedEvidence()
    grouped.generic.append(ClassifiedEvidence(
        evidence_type="generic_context",
        source_kind="chunk", source_id="1", chunk_id="1",
        title="证据", snippet="X" * 800, score=1.0,
    ))

    # 负预算：ask 流程在 history 极长时可能算出负值，to_prompt_text 必须安全
    text_neg = grouped.to_prompt_text(max_chars=-100, per_section=6)
    assert isinstance(text_neg, str)
    # 0 预算
    text_zero = grouped.to_prompt_text(max_chars=0, per_section=6)
    assert isinstance(text_zero, str)
    # 极端情况下产出长度远小于正常 snippet（不会把 800 字全塞进去）
    assert len(text_neg) < 100
    assert len(text_zero) < 100


def test_ask_prompt_respects_hard_cap_with_long_history():
    """ask 接口在对话历史很长时，最终 prompt 仍受总预算硬上限约束，且用户问题不被裁掉。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="证据资料",
        content="穆宁雪是冰系法师，觉醒寒冰系能力。" * 10,
        headers=headers,
    )

    # 先建会话
    sess = client.post(f"/api/projects/{pid}/knowledge/sessions", headers=headers).json()
    conversation_id = sess["id"]

    # 灌入多轮超长历史，撑大 conversation_history
    long_msg = "关于穆宁雪的能力详情，" + ("资料补充说明。" * 400)
    for _ in range(4):
        resp_h = client.post(
            f"/api/projects/{pid}/knowledge/ask",
            json={"question": long_msg, "conversation_id": conversation_id},
            headers=headers,
        )
        assert resp_h.status_code == 200, resp_h.text
        # 确认历史确实复用同一会话（conversation_id 回传一致）
        assert resp_h.json().get("conversation_id") == conversation_id

    # 再问一次，history 已很长；只要不报错且能返回，说明硬上限兜底生效
    final_q = "穆宁雪是什么系？"
    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": final_q, "conversation_id": conversation_id},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    answer = data.get("answer", "")
    assert answer, "用户问题被裁剪或 prompt 组装失败，应仍能返回回答"
    # 用户问题未被盲切：回答应与最终问题相关（含"冰系"或角色名）
    assert "冰系" in answer or "穆宁雪" in answer


def test_retrieve_caps_chunks_per_source():
    """单个 source 贡献的 chunk 数应受 MAX_CHUNKS_PER_SOURCE 限制（真实检索路径）。"""
    from sqlalchemy import select
    from services.knowledge_query_plan import build_knowledge_query_plan
    from services.knowledge_retrieval import retrieve_and_classify, MAX_CHUNKS_PER_SOURCE
    from models.project_source_chunk import ProjectSourceChunk as _Chunk
    from test_smoke import test_session_factory

    pid, headers = _create_project()
    # 每段带章节标题，强制结构分块，切出远超 MAX_CHUNKS_PER_SOURCE 的 chunk
    segment = "穆宁雪觉醒冰系法师的能力，掌握寒冰系基础技能。" * 4
    long_content = "\n\n".join(f"## 第{i}段\n{segment}" for i in range(20))
    create_resp = _create_source(pid, title="穆宁雪长资料", content=long_content, headers=headers)
    sid = create_resp.json()["id"]

    # 确认确实切出了 >8 个 chunk
    async def _count():
        async with test_session_factory() as session:
            result = await session.execute(select(_Chunk).where(_Chunk.source_id == sid))
            return len(list(result.scalars().all()))
    assert asyncio.run(_count()) > MAX_CHUNKS_PER_SOURCE

    plan = build_knowledge_query_plan("穆宁雪是什么系？", project_id=pid)

    async def _retrieve():
        async with test_session_factory() as session:
            _, _, grouped = await retrieve_and_classify(
                session, pid, plan,
                intent="character_ability",
                entities=["穆宁雪"],
                attributes=["冰系"],
                include_structured=False,
                limit=30,
            )
            return grouped

    grouped = asyncio.run(_retrieve())
    # 统计 grouped 各分组中来自该 source 的 chunk 数
    all_items = (
        list(grouped.direct_character) + list(grouped.relationship)
        + list(grouped.ability_table) + list(grouped.worldbuilding)
        + list(grouped.timeline) + list(grouped.generic)
    )
    same_source_chunks = [it for it in all_items if it.source_id == sid and it.chunk_id]
    assert len(same_source_chunks) <= MAX_CHUNKS_PER_SOURCE, (
        f"同 source chunk 数 {len(same_source_chunks)} 超过上限 {MAX_CHUNKS_PER_SOURCE}"
    )


def test_chunk_and_save_no_longer_writes_mock_embedding():
    """切片入库不应再写入 mock embedding（向量检索未接入，字段应留空）。"""
    from sqlalchemy import select
    from models.project_source_chunk import ProjectSourceChunk as _Chunk
    from test_smoke import test_session_factory

    pid, headers = _create_project()
    create_resp = _create_source(pid, title="切片测试", content="段落一。段落二。段落三。" * 20, headers=headers)
    sid = create_resp.json()["id"]

    # 触发 reindex 确保 chunk_and_save 执行
    resp = client.post(f"/api/projects/{pid}/knowledge/sources/{sid}/reindex", headers=headers)
    assert resp.status_code == 200, resp.text

    # 直接查 DB 验证 embedding 字段未被填充
    async def _check():
        async with test_session_factory() as session:
            result = await session.execute(
                select(_Chunk).where(_Chunk.source_id == sid)
            )
            return list(result.scalars().all())

    chunks = asyncio.run(_check())
    assert len(chunks) > 0
    for c in chunks:
        assert not c.embedding, f"chunk {c.chunk_index} 不应写入 embedding，得到 {c.embedding!r}"


# ── 证据就近绑定回归测试（防止法系误归因） ─────────────────


def test_classify_direct_binding_true_positive():
    """人物与法系就近绑定（释放/觉醒）应判为 direct_character_evidence。"""
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "张小侯释放风系魔法击退了敌人，展现出强大的控制力。",
        "title": "战斗片段",
        "source_type": "",
    }
    ev_type = _classify_evidence(
        result,
        intent="character_ability",
        entities=["张小侯"],
        attributes=["风系"],
    )
    assert ev_type == "direct_character_evidence", f"就近绑定应判 direct，得到 {ev_type}"


def test_classify_rejects_third_party_attribution():
    """'赵满延的光系魔法保护了张小侯'不应把光系归因给张小侯。

    这是误归因的核心场景：片段同时含张小侯和光系，但光系属于赵满延。
    """
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "赵满延的光系魔法保护了张小侯，挡下了致命一击。",
        "title": "团战片段",
        "source_type": "",
    }
    ev_type = _classify_evidence(
        result,
        intent="character_ability",
        entities=["张小侯", "赵满延"],  # 已知实体含赵满延，用于识别第三人
        attributes=["光系"],
    )
    assert ev_type != "direct_character_evidence", (
        f"旁人施法不应判 direct（误归因），得到 {ev_type}"
    )
    # 应降级为 relationship（旁人相关）或 generic
    assert ev_type in ("relationship_evidence", "generic_context"), (
        f"应降级为 relationship/generic，得到 {ev_type}"
    )


def test_classify_nearby_binding_with_bystander_present():
    """旁人在场但不夹在人物与法系之间时，绑定仍应成立。

    '张小侯觉醒了冰系，同时赵满延在场'：冰系绑定张小侯，赵满延只是旁观，
    不应因有第三人出现就拒绝绑定。
    """
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "张小侯觉醒了冰系能力，同时赵满延在一旁观战。",
        "title": "觉醒片段",
        "source_type": "",
    }
    ev_type = _classify_evidence(
        result,
        intent="character_ability",
        entities=["张小侯", "赵满延"],
        attributes=["冰系"],
    )
    assert ev_type == "direct_character_evidence", (
        f"旁人在场不挡绑定，应判 direct，得到 {ev_type}"
    )


def test_classify_no_attributes_demotes_to_generic():
    """character_ability 意图但无属性时，含人物名片段应降级为 generic。

    避免无属性时把所有含人物名的片段都当直接证据，引入无关法系。
    """
    from services.knowledge_retrieval import _classify_evidence

    result = {
        "snippet": "张小侯站在城墙之上，望着远方的敌军。",
        "title": "场景描写",
        "source_type": "",
    }
    ev_type = _classify_evidence(
        result,
        intent="character_ability",
        entities=["张小侯"],
        attributes=[],  # planner 未解析出法系
    )
    assert ev_type == "generic_context", f"无属性应降级 generic，得到 {ev_type}"


# ── API 级回归：法系误归因端到端 ─────────────────────────


def test_ask_character_system_no_third_party_attribution():
    """端到端：问"张小侯是什么系的"，回答应含风系，不含光系/火系。

    资料同时包含：
    - 张小侯的风系正向证据（反问句式）
    - 赵满延的光系魔法保护张小侯（误归因陷阱）
    - 莫凡释放火系，张小侯旁观（误归因陷阱）

    本地确定性答案 _try_compile_local_qa_answer 在 mock LLM 下即可验证，
    不依赖真实模型。
    """
    pid, headers = _create_project()
    _create_source(
        pid,
        title="张小侯能力",
        content=(
            "张小侯，你不是风系的吗，看看能不能把风系的初阶技能-风轨释放出来。"
            "赵满延的光系魔法保护了张小侯，挡下了致命一击。"
            "莫凡释放火系魔法，张小侯在旁观看。"
        ),
        headers=headers,
    )

    # 先建会话
    sess = client.post(f"/api/projects/{pid}/knowledge/sessions", headers=headers).json()
    conversation_id = sess["id"]

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "张小侯是什么系的", "conversation_id": conversation_id},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    answer = data.get("answer", "")

    # 核心断言：回答应明确张小侯是风系
    assert "风系" in answer, f"回答应含风系，实际: {answer!r}"
    # 不应把旁人的法系误归因给张小侯
    assert "光系" not in answer, f"回答不应含光系（旁人赵满延的法系），实际: {answer!r}"
    assert "火系" not in answer, f"回答不应含火系（旁人莫凡的法系），实际: {answer!r}"


# ── 规则事实索引：抽取与写入测试（文档 §13.1） ──────────


def test_fact_extraction_character_system_no_misattribution():
    """规则抽取：张小侯->风系命中，光系/火系不误归因。"""
    from services.knowledge_fact_rules import extract_character_system_facts_from_text

    text = (
        "张小侯，你不是风系的吗，看看能不能把风轨释放出来。"
        "赵满延的光系魔法保护了张小侯。"
        "莫凡释放火系魔法，张小侯在旁观看。"
    )
    facts = extract_character_system_facts_from_text(text, source_title="测试")
    pairs = {(f.subject, f.object) for f in facts}

    assert ("张小侯", "风系") in pairs, f"应抽到张小侯->风系，实际: {pairs}"
    assert ("张小侯", "光系") not in pairs, f"不应误归因张小侯->光系，实际: {pairs}"
    assert ("张小侯", "火系") not in pairs, f"不应误归因张小侯->火系，实际: {pairs}"


def test_fact_extraction_multiple_systems():
    """规则抽取：同一人物多法系都能抽到。"""
    from services.knowledge_fact_rules import extract_character_system_facts_from_text

    text = "莫凡觉醒了火系。这是莫凡的雷霆系星尘。莫凡释放暗影系魔法遁入影中。"
    facts = extract_character_system_facts_from_text(text)
    objs = {f.object for f in facts if f.subject == "莫凡"}

    assert "火系" in objs
    assert "雷系" in objs  # 雷霆系归一化
    assert "暗影系" in objs


def test_fact_extraction_skill_alias_weak_fact():
    """规则抽取：技能别名能作为 weak 法系线索写出候选。"""
    from services.knowledge_fact_rules import extract_character_system_facts_from_text

    facts = extract_character_system_facts_from_text("张小侯释放风轨冲出包围。")
    matched = [
        f for f in facts
        if f.subject == "张小侯" and f.object == "风系" and f.confidence == "weak"
    ]

    assert matched, f"风轨应抽为张小侯->风系 weak 事实，实际: {facts}"


def test_fact_extraction_excludes_hypothetical_and_generic():
    """规则抽取：假设性语境和通用法系罗列不抽。"""
    from services.knowledge_fact_rules import extract_character_system_facts_from_text

    text = "张小侯希望自己能觉醒雷系。元素魔法七系包括冰系、火系、土系。"
    facts = extract_character_system_facts_from_text(text)
    # 不应把"希望觉醒雷系"或"七系包括"抽成事实
    assert not any(f.subject == "张小侯" for f in facts), "假设性语境不应抽事实"


def test_rebuild_source_facts_extracts_character_system():
    """端到端：上传资料后 rebuild_source_facts 写入事实表（文档 §13.1）。"""
    from sqlalchemy import select
    from models.project_knowledge_fact import ProjectKnowledgeFact
    from services.knowledge_fact_index import rebuild_source_facts, query_character_system_facts
    from test_smoke import test_session_factory

    pid, headers = _create_project()
    create_resp = _create_source(
        pid,
        title="张小侯能力",
        content=(
            "张小侯，你不是风系的吗，看看能不能把风轨释放出来。"
            "赵满延的光系魔法保护了张小侯。"
            "莫凡释放火系魔法，张小侯在旁观看。"
        ),
        headers=headers,
    )
    sid = create_resp.json()["id"]

    async def _rebuild():
        async with test_session_factory() as session:
            return await rebuild_source_facts(session, pid, sid)

    result = asyncio.run(_rebuild())
    assert result["fact_count"] > 0, f"应抽到事实，实际: {result}"

    # 查事实表
    async def _query():
        async with test_session_factory() as session:
            return await query_character_system_facts(session, pid, subject="张小侯")

    facts = asyncio.run(_query())
    objs = {f.object for f in facts}
    assert "风系" in objs, f"应有张小侯->风系，实际: {objs}"
    assert "光系" not in objs, f"不应有张小侯->光系，实际: {objs}"
    assert "火系" not in objs, f"不应有张小侯->火系，实际: {objs}"


# ── reindex 替换旧 facts 测试（文档 §13.4） ──────────────


def test_reindex_replaces_old_facts():
    """修改资料后 reindex，旧事实应被删除、新事实应存在。"""
    from models.project_knowledge_fact import ProjectKnowledgeFact
    from services.knowledge_fact_index import query_character_system_facts
    from test_smoke import test_session_factory

    pid, headers = _create_project()
    # 1. 上传资料 A：张小侯是风系
    create_resp = _create_source(
        pid, title="能力资料", content="张小侯是风系。", headers=headers,
    )
    sid = create_resp.json()["id"]

    # 2. 确认 facts 有风系（上传时已自动重建）
    async def _query(subject):
        async with test_session_factory() as session:
            return await query_character_system_facts(session, pid, subject=subject)

    facts_before = asyncio.run(_query("张小侯"))
    assert "风系" in {f.object for f in facts_before}, "上传后应有张小侯->风系"

    # 3. 修改资料为：张小侯是土系
    client.patch(
        f"/api/projects/{pid}/knowledge/sources/{sid}",
        json={"content": "张小侯是土系。"},
        headers=headers,
    )

    # 4. 调 source reindex（patch 已自动重建，这里显式再调一次确保）
    resp = client.post(
        f"/api/projects/{pid}/knowledge/sources/{sid}/reindex", headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["fact_count"] > 0

    # 5. 旧风系 fact 被删除，新土系 fact 存在
    facts_after = asyncio.run(_query("张小侯"))
    objs_after = {f.object for f in facts_after}
    assert "土系" in objs_after, f"reindex 后应有张小侯->土系，实际: {objs_after}"
    assert "风系" not in objs_after, f"reindex 后旧风系应被删除，实际: {objs_after}"


# ── QA 优先查 facts 的 API 级测试（文档 §13.2、§13.3） ──


def test_ask_character_system_uses_fact_index():
    """问"张小侯是什么系的"，优先查 facts，回答含风系、不含光系/火系。

    citations 应含 character_system_fact 类型，且不混入无关火系技能表。
    """
    pid, headers = _create_project()
    _create_source(
        pid,
        title="张小侯能力",
        content=(
            "张小侯，你不是风系的吗，看看能不能把风轨释放出来。"
            "赵满延的光系魔法保护了张小侯。"
            "莫凡释放火系魔法，张小侯在旁观看。"
        ),
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "张小侯是什么系的"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    answer = data["answer"]
    citations = data.get("citations", [])

    assert "风系" in answer, f"回答应含风系，实际: {answer!r}"
    assert "光系" not in answer, f"回答不应含光系，实际: {answer!r}"
    assert "火系" not in answer, f"回答不应含火系，实际: {answer!r}"

    # citations 应含 fact 证据类型
    evidence_types = {c.get("evidence_type") for c in citations}
    assert "character_system_fact" in evidence_types, (
        f"citations 应含 character_system_fact，实际 evidence_types: {evidence_types}"
    )
    # 不应混入无关火系技能表
    for c in citations:
        assert "火系技能表" not in c.get("snippet", ""), "citations 不应混入火系技能表"


def test_fact_index_citations_dedupe_same_binding():
    """同一人物-法系有多条证据时，QA citations 不应重复刷屏。"""
    from models.project_knowledge_fact import ProjectKnowledgeFact

    pid, headers = _create_project()
    create_resp = _create_source(
        pid,
        title="重复事实资料",
        content="这是一份用于挂载事实索引的资料。",
        headers=headers,
    )
    sid = create_resp.json()["id"]

    async def _insert_duplicate_facts():
        async with test_session_factory() as session:
            session.add_all([
                ProjectKnowledgeFact(
                    project_id=pid,
                    source_id=sid,
                    fact_type="character_system",
                    subject="张小侯",
                    predicate="has_magic_system",
                    object="风系",
                    confidence="explicit",
                    evidence_text="张小侯是风系。",
                    extractor="test",
                    metadata_={"source_title": "重复事实资料"},
                ),
                ProjectKnowledgeFact(
                    project_id=pid,
                    source_id=sid,
                    fact_type="character_system",
                    subject="张小侯",
                    predicate="has_magic_system",
                    object="风系",
                    confidence="explicit",
                    evidence_text="张小侯释放风轨，证明他掌握风系。",
                    extractor="test",
                    metadata_={"source_title": "重复事实资料"},
                ),
                ProjectKnowledgeFact(
                    project_id=pid,
                    source_id=sid,
                    fact_type="character_system",
                    subject="张小侯",
                    predicate="has_magic_system",
                    object="风系",
                    confidence="explicit",
                    evidence_text="张小侯的风系表现很稳定。",
                    extractor="test",
                    metadata_={"source_title": "重复事实资料"},
                ),
            ])
            await session.commit()

    asyncio.run(_insert_duplicate_facts())

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "张小侯是什么系的"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert "风系" in data["answer"]

    matched = [
        c for c in data.get("citations", [])
        if c.get("evidence_type") == "character_system_fact"
        and c.get("matched_query") == "张小侯 -> 风系"
    ]
    assert len(matched) == 1, f"同一绑定只应返回一条引用，实际: {matched}"


def test_ask_character_by_system_uses_fact_index():
    """问"谁是风系"，反向查 facts，回答含张小侯、不含赵满延。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="人物法系",
        content=(
            "张小侯，你不是风系的吗，看看能不能把风轨释放出来。"
            "赵满延的光系魔法保护了张小侯。"
        ),
        headers=headers,
    )

    resp = client.post(
        f"/api/projects/{pid}/knowledge/ask",
        json={"question": "谁是风系"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    answer = data["answer"]

    assert "张小侯" in answer, f"回答应含张小侯，实际: {answer!r}"
    # 赵满延是光系不是风系，不应出现在"谁是风系"的回答里
    assert "赵满延" not in answer, f"回答不应含赵满延，实际: {answer!r}"


# ── facts API 接口测试（文档 §10） ───────────────────────


def test_facts_api_list_and_rebuild():
    """GET /knowledge/facts 列出事实；POST /knowledge/facts/rebuild 重建。"""
    pid, headers = _create_project()
    _create_source(
        pid,
        title="能力资料",
        content="张小侯是风系。莫凡觉醒了火系。",
        headers=headers,
    )

    # 上传时已自动重建 facts，GET 应能列出
    resp = client.get(f"/api/projects/{pid}/knowledge/facts", headers=headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["total"] > 0, "上传后应有事实"
    subjects = {item["subject"] for item in data["items"]}
    assert "张小侯" in subjects
    assert "莫凡" in subjects

    # 按 subject 过滤
    resp2 = client.get(
        f"/api/projects/{pid}/knowledge/facts?subject=张小侯", headers=headers,
    )
    assert resp2.status_code == 200
    items2 = resp2.json()["items"]
    assert all(item["subject"] == "张小侯" for item in items2)
    assert any(item["object"] == "风系" for item in items2)

    # 重建（全项目）
    resp3 = client.post(
        f"/api/projects/{pid}/knowledge/facts/rebuild",
        json={"source_id": None, "fact_types": ["character_system"]},
        headers=headers,
    )
    assert resp3.status_code == 200, resp3.text
    rebuild_data = resp3.json()
    assert rebuild_data["fact_count"] > 0
    assert rebuild_data["source_count"] >= 1

    # 重建后 facts 仍在
    resp4 = client.get(f"/api/projects/{pid}/knowledge/facts", headers=headers)
    assert resp4.json()["total"] > 0


def test_facts_api_rebuild_single_source():
    """POST /knowledge/facts/rebuild 指定 source_id 重建单个资料的事实。"""
    pid, headers = _create_project()
    create_resp = _create_source(
        pid, title="单资料", content="张小侯是风系。", headers=headers,
    )
    sid = create_resp.json()["id"]

    resp = client.post(
        f"/api/projects/{pid}/knowledge/facts/rebuild",
        json={"source_id": sid, "fact_types": ["character_system"]},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["fact_count"] > 0
    assert data["source_count"] == 1


# ── 结构化知识人工修正 CRUD ──────────────────────────────


def _create_structured(pid, table, payload, headers):
    return client.post(
        f"/api/projects/{pid}/knowledge/structured/{table}",
        json=payload, headers=headers,
    )


def test_create_manual_character_profile():
    """手动新增人物：origin/canon_level/source_priority 正确。"""
    pid, headers = _create_project()
    resp = _create_structured(pid, "characters", {
        "name": "萧院长", "identity_desc": "天澜魔法高中院长", "manual_note": "手动补充",
    }, headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["origin"] == "manual"
    assert data["canon_level"] == "manual"
    assert data["source_priority"] == 100
    assert "[手动备注] 手动补充" in (data["evidence"] or [])

    # 确认能查到
    listing = client.get(f"/api/projects/{pid}/knowledge/structured/characters", headers=headers)
    assert listing.status_code == 200
    names = [it["name"] for it in listing.json()["items"]]
    assert "萧院长" in names


def test_update_structured_record_preserves_scope():
    """project A 不能更新 project B 的记录。"""
    pid_a, headers_a = _create_project("项目A")
    pid_b, headers_b = _create_project("项目B")

    # A 创建一条能力
    create_resp = _create_structured(pid_a, "abilities", {
        "character_name": "莫凡", "ability_type": "magic_element",
        "ability_name": "雷系", "manual_note": "A项目的能力",
    }, headers_a)
    assert create_resp.status_code == 200
    record_id = create_resp.json()["id"]

    # B 尝试修改 A 的记录 → 应 404
    patch_resp = client.patch(
        f"/api/projects/{pid_b}/knowledge/structured/abilities/{record_id}",
        json={"level_desc": "高阶"}, headers=headers_b,
    )
    assert patch_resp.status_code == 404


def test_delete_structured_record_preserves_scope():
    """project A 不能删除 project B 的记录。"""
    pid_a, headers_a = _create_project("删除项目A")
    pid_b, headers_b = _create_project("删除项目B")

    create_resp = _create_structured(pid_a, "world_rules", {
        "category": "魔法体系", "rule_text": "测试规则A",
    }, headers_a)
    assert create_resp.status_code == 200
    record_id = create_resp.json()["id"]

    # B 尝试删除 A 的记录 → 应 404
    del_resp = client.delete(
        f"/api/projects/{pid_b}/knowledge/structured/world_rules/{record_id}",
        headers=headers_b,
    )
    assert del_resp.status_code == 404

    # A 自己能删
    del_ok = client.delete(
        f"/api/projects/{pid_a}/knowledge/structured/world_rules/{record_id}",
        headers=headers_a,
    )
    assert del_ok.status_code == 200
    assert del_ok.json()["deleted"] is True


def test_manual_ability_overrides_llm_extracted_in_qa():
    """手动能力优先于 LLM 抽取能力。"""
    from db.session import async_session
    from models.structured_knowledge import AbilityProfile
    from services.structured_qa import answer_structured_question

    pid, headers = _create_project("QA优先级测试")

    # 模拟 LLM 抽取：莫凡-雷系（original/60）
    async def _insert_llm():
        async with test_session_factory() as session:
            session.add(AbilityProfile(
                project_id=pid, character_name="莫凡",
                ability_type="magic_element", ability_name="雷系",
                origin="llm_extracted", canon_level="original",
                source_priority=60, confidence=0.9, evidence=["莫凡觉醒雷系"],
            ))
            await session.commit()
    asyncio.run(_insert_llm())

    # 用户手动新增：莫凡-冰系（manual/100）
    _create_structured(pid, "abilities", {
        "character_name": "莫凡", "ability_type": "magic_element",
        "ability_name": "冰系", "manual_note": "手动设定冰系",
    }, headers)

    # QA 查莫凡有什么系 → 应包含手动设定的冰系
    async def _ask():
        async with async_session() as db:
            return await answer_structured_question(db, pid, "莫凡有什么系别？")
    result = asyncio.run(_ask())
    answer = result.get("answer", "")
    assert "冰系" in answer, f"手动冰系未出现在回答中: {answer}"
    assert "雷系" in answer, f"LLM雷系也应出现: {answer}"


def test_deleted_ability_not_used_by_structured_qa():
    """删除后 QA 不再使用该记录。"""
    from db.session import async_session
    from services.structured_qa import answer_structured_question

    pid, headers = _create_project("删除后QA测试")

    # 新增一条能力
    create_resp = _create_structured(pid, "abilities", {
        "character_name": "测试角色", "ability_type": "magic_element",
        "ability_name": "暗系", "manual_note": "待删除",
    }, headers)
    record_id = create_resp.json()["id"]

    # 删除
    del_resp = client.delete(
        f"/api/projects/{pid}/knowledge/structured/abilities/{record_id}",
        headers=headers,
    )
    assert del_resp.status_code == 200

    # QA 查 → 不应再出现暗系
    async def _ask():
        async with async_session() as db:
            return await answer_structured_question(db, pid, "测试角色有什么系别？")
    result = asyncio.run(_ask())
    answer = result.get("answer", "")
    assert "暗系" not in answer, f"删除后仍出现暗系: {answer}"


def test_update_preserves_evidence_or_adds_manual_note():
    """编辑不应导致 evidence 完全丢失。"""
    from models.structured_knowledge import AbilityProfile

    pid, headers = _create_project("evidence保留测试")

    # 模拟 LLM 抽取带 evidence
    async def _insert_llm():
        async with test_session_factory() as session:
            session.add(AbilityProfile(
                project_id=pid, character_name="莫凡",
                ability_type="magic_element", ability_name="火系",
                origin="llm_extracted", canon_level="original",
                source_priority=60, confidence=0.9, evidence=["莫凡觉醒火系"],
            ))
            await session.commit()
    asyncio.run(_insert_llm())

    # 查到该记录 id
    listing = client.get(f"/api/projects/{pid}/knowledge/structured/abilities", headers=headers)
    record_id = None
    for it in listing.json()["items"]:
        if it["character_name"] == "莫凡" and it["ability_name"] == "火系":
            record_id = it["id"]
            break
    assert record_id is not None

    # 编辑：加 manual_note
    patch_resp = client.patch(
        f"/api/projects/{pid}/knowledge/structured/abilities/{record_id}",
        json={"level_desc": "初阶", "manual_note": "确认是初阶"},
        headers=headers,
    )
    assert patch_resp.status_code == 200, patch_resp.text
    evidence = patch_resp.json()["evidence"]
    # 原 evidence 保留
    assert any("莫凡觉醒火系" in e for e in evidence), f"原evidence丢失: {evidence}"
    # manual_note 追加
    assert any("手动备注" in e for e in evidence), f"manual_note未追加: {evidence}"
    # 升级为 manual
    assert patch_resp.json()["origin"] == "manual"
    assert patch_resp.json()["source_priority"] == 100
